#!/usr/bin/env python
"""Genera Informe EP2 ISY0101.docx (agente funcional IntegraBot)."""
from __future__ import annotations

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(__file__))

import fill_informe_docx as docx_helpers

Document = docx_helpers.Document
DOC_PATH = os.path.join(ROOT, "Informe EP2 ISY0101.docx")
TEMPLATE_PATH = os.path.join(ROOT, "_informe_work.docx")
FALLBACK_TEMPLATE = os.path.join(ROOT, "Informe EP1 ISY0101.docx")
DIAGRAM_PATH = os.path.join(ROOT, "docs", "diagrama_orquestacion_ep2.png")
COVER_PARAGRAPHS = docx_helpers.COVER_PARAGRAPHS

add_body = docx_helpers.add_body
add_heading = docx_helpers.add_heading
add_bullet = docx_helpers.add_bullet
add_table = docx_helpers.add_table
add_figure = docx_helpers.add_figure
add_index_page = docx_helpers.add_index_page
clear_body_after_cover = docx_helpers.clear_body_after_cover
finalize_cover_page_break = docx_helpers.finalize_cover_page_break
set_paragraph_text = docx_helpers.set_paragraph_text


def update_cover_ep2(doc) -> None:
    mapping = {
        3: "INGENIERÍA DE SOLUCIONES CON INTELIGENCIA ARTIFICIAL",
        4: "ISY0101 — Sección 011V",
        5: "San Bernardo, Chile",
        10: "IntegraBot — Agente funcional de RRHH (EP2)",
        16: "Integrante:",
        18: "Diego Alejandro Madrid Muñoz",
        19: "RUT: 19.914.672-6",
        23: "Prof. Jaime Andrés Menéndez Silva",
        25: "Evaluación Parcial N°2 — Encargo con Presentación",
        26: "17 de octubre de 2026",
    }
    for idx, text in mapping.items():
        p = doc.paragraphs[idx]
        set_paragraph_text(p, text, center=True)
        for run in p.runs:
            docx_helpers.set_run_font(run, bold=(idx == 10))


def add_index_ep2(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
    from docx.shared import Inches, Pt

    p = doc.add_paragraph()
    run = p.add_run("ÍNDICE DE CONTENIDO")
    docx_helpers.set_run_font(run, bold=True)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)

    entries = [
        (0, "Introducción"),
        (0, "1. Caso organizacional y continuidad EP1 → EP2"),
        (0, "2. Herramientas y framework del agente (IE1, IE2)"),
        (0, "3. Memoria y recuperación semántica (IE3, IE4)"),
        (0, "4. Planificación y decisiones adaptativas (IE5)"),
        (0, "5. Orquestación, pruebas y repositorio (IE6, IE8)"),
        (0, "6. Justificación de componentes (IE9)"),
        (0, "Conclusión"),
        (0, "Reflexión individual"),
        (0, "Declaración de uso de inteligencia artificial"),
        (0, "Bibliografía"),
    ]
    for level, entry in entries:
        docx_helpers.add_index_line(doc, entry, level)

    pb = doc.add_paragraph()
    pb.add_run().add_break(WD_BREAK.PAGE)


def build_content_ep2(doc) -> None:
    add_index_ep2(doc)

    add_heading(doc, "Introducción")
    add_body(
        doc,
        "El presente informe documenta la evolución de IntegraBot — iniciado en la "
        "Evaluación Parcial N°1 como copiloto RAG — hacia un agente funcional para "
        "IntegraServicios S.A. (Gestión de Personas). La EP2 exige integrar herramientas "
        "de consulta, escritura y razonamiento; configurar memoria de corto y largo plazo; "
        "demostrar planificación ante tareas multietapa; y entregar documentación técnica "
        "con diagramas, flujos de trabajo y evidencias de prueba en GitHub.",
    )
    add_body(
        doc,
        "Repositorio: https://github.com/DeboradorDeMundos/integraservicios-rrhh. "
        "Ejecución: python generate_data.py && streamlit run app.py. "
        "Pruebas EP2: python scripts/agent_scenarios_test.py [--live].",
    )

    add_heading(doc, "1. Caso organizacional y continuidad EP1 → EP2")
    add_table(
        doc,
        "Tabla 1. Objetivos de la EP2",
        ["Objetivo", "Meta"],
        [
            ["Autonomía controlada", "Elegir búsqueda, memoria o ticket según el caso"],
            ["Continuidad conversacional", "Memoria corta y larga por sesión"],
            ["Trazabilidad", "Plan, tools, fuentes RAG y ticket INT visibles"],
            ["Herencia EP1", "Corpus dual y guardrails anti-alucinación"],
        ],
    )
    add_table(
        doc,
        "Tabla 2. Corpus documental",
        ["Tipo", "Documento", "Tool asociada"],
        [
            ["Interno", "RIOHS_IntegraServicios_2026.pdf", "buscar_politicas_internas"],
            ["Interno", "Contrato_Colectivo_2026.pdf", "buscar_politicas_internas"],
            ["Externo", "Codigo_del_Trabajo_Chile.pdf", "buscar_normativa_externa"],
        ],
    )

    add_heading(doc, "2. Herramientas y framework del agente (IE1, IE2)")
    add_body(
        doc,
        "El agente se implementa con LangChain (mensajes, tools, bind_tools) y un bucle "
        "ReAct de hasta cuatro iteraciones (src/agent.py), patrón equivalente a LangChain "
        "Agents. Cuatro herramientas ejecutan funciones específicas con autonomía: "
        "búsqueda en corpus interno, búsqueda en normativa externa, persistencia de "
        "memoria de caso y registro de tickets RRHH mock (INT-xxxxx).",
    )
    add_table(
        doc,
        "Tabla 3. Herramientas del agente IntegraBot",
        ["Herramienta", "Tipo", "Función"],
        [
            ["buscar_politicas_internas", "Consulta", "RIOHS + Contrato Colectivo (Chroma filter interno)"],
            ["buscar_normativa_externa", "Consulta", "Código del Trabajo / Ley Karin (filter externo)"],
            ["guardar_memoria_caso", "Escritura", "Hechos y tema en ./case_memory/{session}.json"],
            ["registrar_ticket_gestion_personas", "Escritura", "Ticket en tickets_local/tickets_rrhh.jsonl"],
        ],
    )
    add_table(
        doc,
        "Tabla 4. Stack tecnológico del agente",
        ["Capa", "Tecnología", "Ubicación"],
        [
            ["Orquestación", "LangChain ReAct, bind_tools", "src/agent.py"],
            ["Planificación", "Clasificación + plan numerado", "src/planner.py"],
            ["RAG", "ChromaDB, MMR k=4", "src/retriever.py, src/rag_engine.py"],
            ["Interfaz", "Streamlit + expander orquestación", "app.py"],
        ],
    )
    add_body(
        doc,
        "El RAGPipeline EP1 se reutiliza como motor documental; el agente añade decisión y "
        "escritura. MODEL_PROVIDER google/groq usa tool-calling; local ejecuta el plan de "
        "forma determinista para demos sin API.",
    )

    add_heading(doc, "3. Memoria y recuperación semántica (IE3, IE4)")
    add_body(
        doc,
        "Memoria de corto plazo: AGENT_HISTORY_LIMIT=8 turnos inyectados vía conversation.py. "
        "Memoria de largo plazo: CaseMemoryStore persiste colaborador, tema, hechos y tickets "
        "por session_id de Streamlit, permitiendo follow-ups («¿y el ticket que abrimos?»).",
    )
    add_body(
        doc,
        "Recuperación semántica: retriever.py centraliza ChromaDB con MMR k=4 y filtros "
        "metadata source_type (interno/externo), asegurando continuidad en flujos prolongados "
        "sin mezclar políticas corporativas con normativa externa.",
    )
    add_table(
        doc,
        "Tabla 5. Capas de memoria y contexto",
        ["Capa", "Módulo", "Parámetro / propósito"],
        [
            ["Corto plazo", "conversation.py", "AGENT_HISTORY_LIMIT=8 turnos"],
            ["Largo plazo", "memory.py", "JSON por session_id (hechos, tickets)"],
            ["Semántico", "retriever.py + ChromaDB", "MMR k=4, filtro source_type"],
        ],
    )
    add_table(
        doc,
        "Tabla 6. Parámetros de recuperación (herencia EP1)",
        ["Parámetro", "Valor", "Justificación"],
        [
            ["Retriever", "MMR", "Balance relevancia y diversidad"],
            ["k", "4", "Trazabilidad en expander de fuentes"],
            ["Filtro", "source_type interno/externo", "Separar reglamento y ley"],
        ],
    )

    add_heading(doc, "4. Planificación y decisiones adaptativas (IE5)")
    add_body(
        doc,
        "src/planner.py clasifica cada consulta (FAQ, multietapa, sensible, fuera de dominio, "
        "seguimiento) y genera un plan numerado con prioridades. El plan se muestra en la UI "
        "(expander Orquestación del agente) y se inyecta al System Prompt del agente.",
    )
    add_table(
        doc,
        "Tabla 7. Escenarios de demostración",
        ["#", "Pregunta / escenario", "Plan", "Evidencia"],
        [
            ["1", "Bono de colación", "FAQ", "Cita + fuentes RAG"],
            ["2", "Días admin. y documentos", "Multietapa", "Plan + memoria"],
            ["3", "¿Mañana lloverá?", "Fuera de dominio", "Derivación sin RAG"],
            ["4", "Finiquito / cheque litigioso", "Sensible", "Ticket INT + escalamiento"],
            ["5", "Estado del ticket abierto", "Seguimiento", "Memoria + estado ticket"],
        ],
    )

    add_heading(doc, "5. Orquestación, pruebas y repositorio (IE6, IE8)")
    add_body(
        doc,
        "La orquestación conecta app.py → run_agent → planner → tools → RAG/guardrails → "
        "respuesta con trazabilidad. scripts/agent_scenarios_test.py valida planificación, "
        "memoria persistente y tres modos de decisión; scripts/smoke_test.py mantiene regresión EP1.",
    )
    add_table(
        doc,
        "Tabla 8. Módulos y responsabilidades",
        ["Módulo", "Responsabilidad"],
        [
            ["app.py", "Chat Streamlit, expander orquestación"],
            ["agent.py", "run_agent, bucle ReAct"],
            ["planner.py", "Clasificación y plan de ejecución"],
            ["agent_tools.py", "Cuatro tools consulta/escritura"],
            ["memory.py", "Persistencia de caso por sesión"],
            ["guardrails.py", "Validación post-respuesta (EP1)"],
        ],
    )
    add_figure(
        doc,
        DIAGRAM_PATH,
        "Figura 1. Diagrama de orquestación del agente IntegraBot (EP2).",
    )
    add_table(
        doc,
        "Tabla 9. Scripts de prueba",
        ["Script", "Modo", "Validación"],
        [
            ["scripts/smoke_test.py", "Offline / --live", "Regresión corpus RAG EP1"],
            ["scripts/agent_scenarios_test.py", "Offline", "Planner, memoria, sensible"],
            ["streamlit run app.py", "Manual", "Demo integral y capturas AVA"],
        ],
    )
    add_table(
        doc,
        "Tabla 10. Checklist de entrega AVA",
        ["Paso", "Acción"],
        [
            ["1", "Informe Word/PDF (máx. 5 páginas de cuerpo)"],
            ["2", "Capturas: orquestación, fuentes, ticket INT"],
            ["3", "Subir AVA + enlace GitHub + correo docente"],
            ["4", "Presentación oral con demo en vivo"],
        ],
    )

    add_heading(doc, "6. Justificación de componentes (IE9)")
    add_table(
        doc,
        "Tabla 11. Decisiones de diseño",
        ["Decisión", "Motivo"],
        [
            ["ReAct + tools vs. cadena fija EP1", "EP2 exige autonomía con consulta y escritura"],
            ["Dos retrievers filtrados", "Separar política interna y normativa externa"],
            ["Memoria JSON por sesión", "Flujos multietapa sin inflar tokens"],
            ["Tickets mock INT", "Escritura auditable y reproducible"],
            ["Guardrails EP1", "No relajar cumplimiento por tener agente"],
        ],
    )

    add_heading(doc, "Conclusión")
    add_body(
        doc,
        "[Completar en tus palabras, sin IA: aporte del agente a IntegraServicios, "
        "relación con EP1 y aprendizajes principales.]",
    )

    add_heading(doc, "Reflexión individual", level=2)
    add_body(
        doc,
        "Diego Alejandro Madrid Muñoz: [8–12 líneas sin IA: tu rol, dificultades "
        "(API, casos sensibles, demo) y aprendizaje.]",
    )

    add_heading(doc, "Declaración de uso de inteligencia artificial", level=2)
    add_body(
        doc,
        "Se empleó IA generativa (Cursor IDE, Google Gemini API) como apoyo en revisión de "
        "código, borrador de documentación, diagramas y presentación HTML. El análisis del "
        "caso, las decisiones de diseño, las conclusiones y la reflexión individual son "
        "responsabilidad del autor. Duoc UC (2025): https://bibliotecas.duoc.cl/ia",
    )

    add_heading(doc, "Bibliografía")
    refs = [
        "Chroma. (2026). ChromaDB documentation. https://docs.trychroma.com/",
        "Duoc UC. (2025). Uso de inteligencia artificial en el aprendizaje. https://bibliotecas.duoc.cl/ia",
        "Google. (2026). Gemini API documentation. https://ai.google.dev/",
        "LangChain. (2026). LangChain Python documentation. https://python.langchain.com/",
        "Streamlit. (2026). Streamlit documentation. https://docs.streamlit.io/",
    ]
    for ref in refs:
        docx_helpers.add_body(doc, ref)


def main() -> None:
    diagram_script = os.path.join(ROOT, "scripts", "render_diagram_ep2.py")
    if os.path.isfile(diagram_script):
        import subprocess

        subprocess.run([sys.executable, diagram_script], check=False)

    source = TEMPLATE_PATH if os.path.isfile(TEMPLATE_PATH) else FALLBACK_TEMPLATE
    if not os.path.isfile(source):
        print(f"No hay plantilla Word: {source}", file=sys.stderr)
        raise SystemExit(1)

    doc = Document(source)
    update_cover_ep2(doc)
    finalize_cover_page_break(doc)
    clear_body_after_cover(doc)
    build_content_ep2(doc)
    try:
        doc.save(DOC_PATH)
        print(f"Informe EP2 generado: {DOC_PATH}")
    except PermissionError:
        alt = os.path.join(ROOT, "Informe EP2 ISY0101_corregido.docx")
        doc.save(alt)
        print(f"Cierra Word. Guardado en: {alt}", file=sys.stderr)


if __name__ == "__main__":
    main()
